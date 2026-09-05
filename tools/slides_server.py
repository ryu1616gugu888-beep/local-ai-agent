#!/usr/bin/env python3
"""PowerPoint(.pptx)スライド生成用の自作MCPサーバー(python-pptx使用、無料・オフライン)。

各スライドはタイトル・本文の箇条書き・画像(任意、generate_imageで生成したパス等)を
指定できる。生成したファイルは data/slides/ に保存される。
"""

import time
from pathlib import Path

from mcp.server.mcpserver import MCPServer
from pptx import Presentation
from pptx.util import Inches

OUTPUT_DIR = Path(__file__).parent.parent / "data" / "slides"

mcp = MCPServer("slides")


@mcp.tool()
def create_presentation(title: str, slides: list) -> str:
    """タイトルとスライド一覧からPowerPointファイル(.pptx)を作成する。

    slides は各要素が {"title": str, "bullets": [str, ...], "image_path": str(任意)}
    という形式のリスト。image_path は generate_image ツールが返したファイルパスなど、
    このMac上に実在する画像ファイルを指定する。作成したファイルのパスを返す。
    """
    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title

    for spec in slides:
        layout = prs.slide_layouts[1]  # タイトル + コンテンツ
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = spec.get("title", "")

        bullets = spec.get("bullets") or []
        if bullets:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = bullets[0]
            for b in bullets[1:]:
                p = tf.add_paragraph()
                p.text = b

        image_path = spec.get("image_path")
        if image_path and Path(image_path).exists():
            slide.shapes.add_picture(image_path, Inches(5.5), Inches(1.5), height=Inches(4))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    safe_title = "".join(c for c in title if c.isalnum() or c in " _-")[:40].strip() or "slides"
    dest = OUTPUT_DIR / f"{safe_title}_{int(time.time())}.pptx"
    prs.save(str(dest))
    return f"スライドを作成しました: {dest}"


if __name__ == "__main__":
    mcp.run()
